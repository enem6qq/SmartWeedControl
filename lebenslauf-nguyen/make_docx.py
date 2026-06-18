#!/usr/bin/env python3
"""Erzeugt eine bearbeitbare Word-Version des Lebenslaufs.

Layout: zweispaltige Tabelle (hellblaue Sidebar links, Inhalt rechts),
Petrol-Akzentfarbe, Schrift "Blinker" – passend zur HTML/PDF-Vorlage.

Aufruf:  python3 make_docx.py
Ausgabe: Lebenslauf_Trang_Nguyen.docx
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).parent
PHOTO = HERE / "assets" / "img" / "foto.png"

# ---- Farben (aus dem Original-Word-Dokument) ----
INK        = RGBColor(0x2A, 0x2A, 0x2A)
ACCENT     = RGBColor(0x10, 0x65, 0x7E)
ACCENT_DK  = RGBColor(0x0B, 0x4A, 0x5C)
MUTED      = RGBColor(0x4C, 0x60, 0x68)
SIDEBAR    = "CFE0E5"
FONT       = "Blinker"

SIDEBAR_W = Cm(6.6)
MAIN_W    = Cm(14.4)


def set_cell_bg(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(cell, top, left, bottom, right):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        e = OxmlElement(f"w:{tag}")
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)


def no_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "none")
        e.set(qn("w:sz"), "0")
        borders.append(e)
    tblPr.append(borders)


def bottom_border(paragraph, color="10657E", size=14):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    # pBdr muss laut OOXML-Schema vor tabs/spacing/ind stehen
    anchor = None
    for tag in ("w:tabs", "w:spacing", "w:ind", "w:jc", "w:rPr"):
        anchor = pPr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(pbdr)
    else:
        pPr.append(pbdr)


def style_run(run, size=10.5, bold=False, color=INK, caps=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if caps:
        run.font.all_caps = True
    # auch fuer Nicht-Latin-Bereiche dieselbe Schrift
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attr), FONT)


def p(cell, space_before=0, space_after=2, line=1.15):
    para = cell.add_paragraph()
    pf = para.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    return para


def sidebar_heading(cell, text):
    para = p(cell, space_before=10, space_after=4)
    r = para.add_run(text)
    style_run(r, size=11, bold=True, color=ACCENT, caps=True)
    bottom_border(para)


def main_heading(cell, text):
    para = p(cell, space_before=12, space_after=6)
    r = para.add_run(text)
    style_run(r, size=13, bold=True, color=ACCENT, caps=True)
    bottom_border(para, size=18)


def bullet(cell, text, color=INK, size=9.8):
    para = p(cell, space_after=2)
    para.paragraph_format.left_indent = Cm(0.35)
    r = para.add_run("•  ")
    style_run(r, size=size, color=ACCENT)
    r2 = para.add_run(text)
    style_run(r2, size=size, color=color)
    return para


def build():
    doc = Document()

    # Seitenraender auf 0 -> Sidebar reicht bis zum Rand
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    for side in ("top", "bottom", "left", "right"):
        setattr(sec, f"{side}_margin", Cm(0))
    sec.header_distance = Cm(0)
    sec.footer_distance = Cm(0)

    # Standard-Schrift
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.allow_autofit = False
    no_table_borders(table)

    # volle Seitenhoehe erzwingen (Sidebar-Faerbung)
    trPr = table.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), "15300")
    h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)

    left, right = table.rows[0].cells
    left.width = SIDEBAR_W
    right.width = MAIN_W
    set_cell_bg(left, SIDEBAR)
    set_cell_margins(left, 480, 360, 480, 360)
    set_cell_margins(right, 540, 420, 480, 480)

    # erste (leere) Standardabsaetze entfernen
    for c in (left, right):
        c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p)

    # ---------------- SIDEBAR ----------------
    if PHOTO.exists():
        pic_p = left.add_paragraph()
        pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_p.paragraph_format.space_after = Pt(6)
        pic_p.add_run().add_picture(str(PHOTO), width=Cm(4.6))

    sidebar_heading(left, "Kontakt")
    for line in ("Wörthstraße 66", "72766 Reutlingen"):
        para = p(left, space_after=0); style_run(para.add_run(line), size=9.6)
    para = p(left, space_before=4); style_run(para.add_run("0152 22 71 68 19"), size=9.6)
    para = p(left, space_before=4); style_run(para.add_run("trangnguyen240591@gmail.com"), size=9.6)

    sidebar_heading(left, "Persönliches")
    persoenlich = [
        ("Geburtsdatum", "24.05.1991"),
        ("Geburtsort", "Reutlingen"),
        ("Staatsangehörigkeit", "Deutsch"),
        ("Familienstand", "Ledig, ein Kind"),
    ]
    for label, value in persoenlich:
        para = p(left, space_before=4, space_after=0)
        style_run(para.add_run(label), size=8.4, bold=True, color=ACCENT_DK, caps=True)
        para2 = p(left, space_after=0)
        style_run(para2.add_run(value), size=9.6)

    sidebar_heading(left, "Qualifikationen")
    bullet(left, "Führerschein Klasse B")

    sidebar_heading(left, "Sprachen")
    for lang, lvl in (("Deutsch", "Muttersprache"), ("Vietnamesisch", "Muttersprache")):
        para = p(left, space_after=1)
        style_run(para.add_run(lang), size=9.8, bold=True, color=INK)
        para2 = p(left, space_after=3)
        style_run(para2.add_run(lvl), size=9, color=MUTED)

    # ---------------- HAUPTSPALTE ----------------
    name_p = p(right, space_after=0)
    style_run(name_p.add_run("Trang Phuong Nguyen"), size=26, bold=True, color=ACCENT)
    role_p = p(right, space_before=2, space_after=4)
    style_run(role_p.add_run("Bäckereifachverkäuferin"), size=11.5, bold=True, color=MUTED)

    def entry(title, date, org):
        para = p(right, space_before=6, space_after=0)
        para.paragraph_format.tab_stops.add_tab_stop(Cm(12.6), WD_TAB_ALIGNMENT.RIGHT)
        style_run(para.add_run(title), size=11, bold=True, color=INK)
        style_run(para.add_run("\t" + date), size=9.4, bold=True, color=ACCENT)
        org_p = p(right, space_after=2)
        style_run(org_p.add_run(org), size=10, bold=True, color=MUTED)

    main_heading(right, "Beruflicher Werdegang")
    jobs = [
        ("Bäckereifachverkäuferin (Vollzeit)", "03/2011 – heute", "Bio-Vollkornbäckerei Berger, Reutlingen"),
        ("Bäckereifachverkäuferin (Vollzeit)", "11/2010 – 02/2011", "Bäckerei Keim, Reutlingen"),
        ("Bäckereifachverkäuferin (Teilzeit, „Springer“)", "08/2010 – 11/2010", "K&U Bäckerei"),
        ("Bäckereifachverkäuferin (Vollzeit)", "07/2009 – 07/2010", "Hotel Garni (Schwanen), Reutlingen-Betzingen"),
        ("Ausbildung zur Bäckereifachverkäuferin", "09/2006 – 07/2009", "Hotel Garni (Schwanen), Reutlingen-Betzingen"),
    ]
    for t, d, o in jobs:
        entry(t, d, o)

    main_heading(right, "Schulischer Werdegang")
    schools = [
        ("Kerschensteiner Schule (Berufsschule)", "09/2006 – 07/2009", "Reutlingen"),
        ("Eduard-Spranger-Schule (Hauptschulabschluss)", "09/2002 – 07/2006", "Reutlingen"),
        ("Eduard-Spranger-Schule (Grundschule)", "09/1997 – 07/2002", "Reutlingen"),
    ]
    for t, d, o in schools:
        entry(t, d, o)

    sig = p(right, space_before=16)
    style_run(sig.add_run("Reutlingen, 18.06.2026"), size=9.8, color=MUTED)

    out = HERE / "Lebenslauf_Trang_Nguyen.docx"
    doc.save(str(out))
    print("Fertig ->", out.name)


if __name__ == "__main__":
    build()
